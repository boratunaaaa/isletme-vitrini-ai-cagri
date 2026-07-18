"""Lead içe aktarma: XLSX / CSV / JSON.

Akış: dosya → satırlar → kolon tahmini → kullanıcı eşleştirmesi → önizleme →
normalize + dedup → import raporu. Telefonlar normalize edilip şifrelenir;
düz metin telefon rapora/loga yazılmaz (örnek hatalarda maskelenir).
"""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any, BinaryIO

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.crypto import encrypt_pii, hash_phone
from app.core.phone import PhoneValidationError, normalize_phone
from app.models.enums import LeadSource, LeadStatus
from app.models.lead import ImportJob, Lead

TARGET_FIELDS = [
    "business_name",
    "sector",
    "phone",
    "city",
    "district",
    "address",
    "website",
    "external_place_id",
    "demo_link",
    "status",
    "notes",
    "rating",
    "review_count",
]

# Türkçe/İngilizce başlık kalıpları → hedef alan
_HEADER_PATTERNS: list[tuple[str, str]] = [
    (r"firma|i[sş]letme|business|company|unvan|ad[ıi]?$|^ad$|name", "business_name"),
    (r"sekt[oö]r|kategori|category|t[uü]r", "sector"),
    (r"telefon|phone|gsm|cep|numara|tel\b", "phone"),
    (r"[sş]ehir|city|^il$|province", "city"),
    (r"il[cç]e|district|semt|b[oö]lge", "district"),
    (r"adres|address", "address"),
    (r"web|site|url|domain", "website"),
    (r"place.?id|google.?id", "external_place_id"),
    (r"demo", "demo_link"),
    (r"durum|status", "status"),
    (r"not|note|a[cç][ıi]klama|comment", "notes"),
    (r"puan|rating|y[ıi]ld[ıi]z", "rating"),
    (r"yorum|review", "review_count"),
]


def _normalize_header(header: str) -> str:
    text = unicodedata.normalize("NFKC", str(header)).strip().lower()
    return text


def guess_column_mapping(headers: list[str]) -> dict[str, str]:
    """{hedef_alan: kolon_adı} tahmini. İlk eşleşen kolon kazanır."""
    mapping: dict[str, str] = {}
    for header in headers:
        normalized = _normalize_header(header)
        for pattern, target in _HEADER_PATTERNS:
            if target not in mapping and re.search(pattern, normalized):
                mapping[target] = header
                break
    return mapping


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[dict[str, Any]]
    sheet_names: list[str] = field(default_factory=list)


def parse_file(
    file: BinaryIO, filename: str, sheet_name: str | None = None
) -> ParsedFile:
    lower = filename.lower()
    raw = file.read()
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return _parse_xlsx(io.BytesIO(raw), sheet_name)
    if lower.endswith(".csv"):
        return _parse_csv(raw)
    if lower.endswith(".json"):
        return _parse_json(raw)
    raise ValueError("Desteklenmeyen dosya türü (xlsx, csv veya json olmalı)")


def _parse_xlsx(stream: io.BytesIO, sheet_name: str | None) -> ParsedFile:
    workbook = load_workbook(stream, read_only=True, data_only=True)
    sheet_names = list(workbook.sheetnames)
    sheet = workbook[sheet_name] if sheet_name and sheet_name in sheet_names else workbook.active
    if sheet is None:
        raise ValueError("Excel dosyasında sayfa bulunamadı")
    iterator = sheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as exc:
        raise ValueError("Dosya boş") from exc
    headers = [str(h).strip() if h is not None else f"Kolon{i+1}" for i, h in enumerate(header_row)]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        rows.append(
            {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        )
    workbook.close()
    return ParsedFile(headers=headers, rows=rows, sheet_names=sheet_names)


def _parse_csv(raw: bytes) -> ParsedFile:
    for encoding in ("utf-8-sig", "utf-8", "iso-8859-9"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("CSV kodlaması çözülemedi (UTF-8 veya ISO-8859-9 olmalı)")
    sample = text[:4096]
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = list(reader.fieldnames or [])
    rows = [dict(row) for row in reader if any((v or "").strip() for v in row.values())]
    return ParsedFile(headers=headers, rows=rows)


def _parse_json(raw: bytes) -> ParsedFile:
    data = json.loads(raw.decode("utf-8"))
    if isinstance(data, dict):
        for key in ("leads", "rows", "data", "items"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            raise ValueError("JSON bir liste veya {leads: [...]} yapısında olmalı")
    if not isinstance(data, list) or not data:
        raise ValueError("JSON boş veya liste değil")
    headers_set: dict[str, None] = {}
    rows: list[dict[str, Any]] = []
    for item in data:
        if isinstance(item, dict):
            rows.append(item)
            for k in item:
                headers_set.setdefault(str(k))
    return ParsedFile(headers=list(headers_set), rows=rows)


# ─── Commit ──────────────────────────────────────────────────────────────────


@dataclass
class ImportResult:
    job_id: str
    total_rows: int
    imported: int
    duplicates: int
    invalid: int
    errors: list[dict[str, Any]]


_SOURCE_BY_EXT = {
    "xlsx": LeadSource.XLSX_IMPORT.value,
    "csv": LeadSource.CSV_IMPORT.value,
    "json": LeadSource.JSON_IMPORT.value,
}


def _cell(row: dict[str, Any], mapping: dict[str, str], target: str) -> str | None:
    column = mapping.get(target)
    if not column:
        return None
    value = row.get(column)
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def commit_import(
    db: Session,
    *,
    parsed_rows: list[dict[str, Any]],
    mapping: dict[str, str],
    filename: str,
    created_by: uuid.UUID | None,
) -> ImportResult:
    if "business_name" not in mapping:
        raise ValueError("Firma adı kolonu eşleştirilmeden içe aktarma yapılamaz")

    ext = filename.rsplit(".", 1)[-1].lower()
    source = _SOURCE_BY_EXT.get(ext, LeadSource.MANUAL.value)

    existing_hashes: set[str] = {
        h
        for h in db.scalars(select(Lead.phone_hash).where(Lead.phone_hash.is_not(None))).all()
        if h is not None
    }
    existing_place_ids: set[str] = {
        p
        for p in db.scalars(
            select(Lead.external_place_id).where(Lead.external_place_id.is_not(None))
        ).all()
        if p is not None
    }

    imported = duplicates = invalid = 0
    errors: list[dict[str, Any]] = []
    seen_hashes: set[str] = set()
    seen_place_ids: set[str] = set()
    now = datetime.now(UTC)

    for index, row in enumerate(parsed_rows, start=2):  # başlık 1. satır
        business_name = _cell(row, mapping, "business_name")
        if not business_name:
            invalid += 1
            if len(errors) < 50:
                errors.append({"row": index, "error": "Firma adı boş"})
            continue

        place_id = _cell(row, mapping, "external_place_id")
        phone_raw = _cell(row, mapping, "phone")

        phone_fields: dict[str, Any] = {}
        phone_hash_value: str | None = None
        if phone_raw:
            try:
                normalized = normalize_phone(phone_raw)
                phone_hash_value = hash_phone(normalized.e164)
                phone_fields = {
                    "phone_ciphertext": encrypt_pii(normalized.e164),
                    "phone_hash": phone_hash_value,
                    "phone_last_four": normalized.last_four,
                    "phone_country_code": normalized.country_code,
                }
            except PhoneValidationError:
                invalid += 1
                if len(errors) < 50:
                    errors.append(
                        {"row": index, "error": "Geçersiz telefon", "firma": business_name[:60]}
                    )
                continue

        is_duplicate = (
            (
                phone_hash_value
                and (phone_hash_value in existing_hashes or phone_hash_value in seen_hashes)
            )
            or (place_id and (place_id in existing_place_ids or place_id in seen_place_ids))
        )
        if is_duplicate:
            duplicates += 1
            continue

        rating_raw = _cell(row, mapping, "rating")
        review_raw = _cell(row, mapping, "review_count")
        try:
            rating = float(str(rating_raw).replace(",", ".")) if rating_raw else None
        except ValueError:
            rating = None
        try:
            review_count = int(float(str(review_raw).replace(",", "."))) if review_raw else None
        except ValueError:
            review_count = None

        lead = Lead(
            external_place_id=place_id,
            source=source,
            business_name=business_name[:500],
            sector=_cell(row, mapping, "sector"),
            city=_cell(row, mapping, "city"),
            district=_cell(row, mapping, "district"),
            address=_cell(row, mapping, "address"),
            website=_cell(row, mapping, "website"),
            demo_link=_cell(row, mapping, "demo_link"),
            rating=rating,
            review_count=review_count,
            status=LeadStatus.NEW.value,
            imported_at=now,
            **phone_fields,
        )
        db.add(lead)
        if phone_hash_value:
            seen_hashes.add(phone_hash_value)
        if place_id:
            seen_place_ids.add(place_id)
        imported += 1

    job = ImportJob(
        filename=filename[:500],
        source_format=ext,
        status="completed",
        total_rows=len(parsed_rows),
        imported_count=imported,
        duplicate_count=duplicates,
        invalid_count=invalid,
        column_mapping=mapping,
        error_sample=errors,
        created_by=created_by,
    )
    db.add(job)
    db.flush()
    return ImportResult(
        job_id=str(job.id),
        total_rows=len(parsed_rows),
        imported=imported,
        duplicates=duplicates,
        invalid=invalid,
        errors=errors,
    )
